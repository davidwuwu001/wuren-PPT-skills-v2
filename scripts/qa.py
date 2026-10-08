#!/usr/bin/env python3
"""Wuren PPT visual QA. See references/visual-qa-checklist.md for details."""
import sys
import os
import argparse
from pptx import Presentation

SLIDE_W_IN = 13.33
SLIDE_H_IN = 7.5
LOGO_ZONE = (11.4, 0.2, 12.95, 0.75)
PAGENUM_ZONE = (11.8, 7.0, 13.0, 7.4)
FOOTER_ZONE = (0.4, 7.0, 9.0, 7.4)
RATIO_TOLERANCE = 0.05


def emu_to_in(emu):
    return emu / 914400.0 if emu else 0


def rect_overlap(a, b):
    x1 = max(a[0], b[0]); y1 = max(a[1], b[1])
    x2 = min(a[2], b[2]); y2 = min(a[3], b[3])
    if x2 <= x1 or y2 <= y1: return 0
    return (x2 - x1) * (y2 - y1)


def check_deck(pptx_path):
    warnings = []
    if not os.path.exists(pptx_path):
        return [f"File not found: {pptx_path}"]

    pres = Presentation(pptx_path)
    n = len(pres.slides)
    print(f"=== {os.path.basename(pptx_path)} ===")
    print(f"  {n} slides, {pres.slide_width/914400:.2f} x {pres.slide_height/914400:.2f} inch")

    tangzai_streak = 0
    for i, slide in enumerate(pres.slides, 1):
        slide_warns = []
        text_shapes = []
        image_shapes = []
        for sh in slide.shapes:
            try:
                x = emu_to_in(sh.left); y = emu_to_in(sh.top)
                w = emu_to_in(sh.width); h = emu_to_in(sh.height)
                rect = (x, y, x + w, y + h)
            except Exception:
                continue

            if sh.has_text_frame:
                t = sh.text_frame.text
                if not t.strip(): continue
                has_face = any(run.font.name for para in sh.text_frame.paragraphs for run in para.runs)
                if not has_face and len(t) > 5:
                    slide_warns.append(f"  text missing fontFace: {t[:30]!r}")
                if x + w > SLIDE_W_IN + 0.05 or y + h > SLIDE_H_IN + 0.05 or x < -0.05 or y < -0.05:
                    slide_warns.append(f"  text overflows edge: {t[:30]!r} at ({x:.2f},{y:.2f})")
                # Skip system elements that legitimately live in these zones
                t_stripped = t.strip().replace(" ", "")
                is_system = (
                    "/" in t_stripped and t_stripped.replace("/", "").isdigit()  # page nums like "2/15"
                    or t.startswith("PART ")  # PART labels
                    or (t.startswith("五仁汤圆家") and len(t) < 30)
                    or "让运动回归" in t
                )
                if not is_system:
                    for zone, name in [(LOGO_ZONE, "logo"), (PAGENUM_ZONE, "page-num"), (FOOTER_ZONE, "footer")]:
                        if rect_overlap(rect, zone) > 0.05:
                            slide_warns.append(f"  text overlaps {name} zone: {t[:30]!r}")
                text_shapes.append((rect, t))

            if sh.shape_type == 13:
                image_shapes.append((rect, sh))

        for rect, sh in image_shapes:
            disp_w = rect[2] - rect[0]; disp_h = rect[3] - rect[1]
            disp_ratio = disp_w / disp_h if disp_h else 0
            try:
                from PIL import Image
                import io
                im = Image.open(io.BytesIO(sh.image.blob))
                src_ratio = im.size[0] / im.size[1]
                if abs(disp_ratio - src_ratio) > RATIO_TOLERANCE:
                    if src_ratio < 1 and disp_ratio > src_ratio + RATIO_TOLERANCE:
                        slide_warns.append(f"  image stretched: display {disp_ratio:.2f} vs source {src_ratio:.2f}")
                    elif src_ratio > 1 and disp_ratio < src_ratio - RATIO_TOLERANCE:
                        slide_warns.append(f"  image stretched: display {disp_ratio:.2f} vs source {src_ratio:.2f}")
            except Exception:
                pass

        has_tangzai = False
        for sh in slide.shapes:
            if sh.shape_type == 13:
                try:
                    if sh.image.filename and 'tangzai' in sh.image.filename.lower():
                        has_tangzai = True
                except Exception: pass
        if has_tangzai:
            tangzai_streak += 1
            if tangzai_streak >= 3:
                slide_warns.append(f"  tangzai used {tangzai_streak} consecutive slides")
        else:
            tangzai_streak = 0

        for rect, t in text_shapes:
            if rect[1] < 1.0 and (rect[3] - rect[1]) > 0.5 and len(t) > 4:
                title_bottom = rect[3]
                for sh in slide.shapes:
                    if sh.shape_type == 1:
                        try:
                            sy = emu_to_in(sh.top); sh_h = emu_to_in(sh.height)
                        except Exception: continue
                        if sy < title_bottom + 0.05 and sy + sh_h > title_bottom - 0.1 and sh_h < 0.15:
                            slide_warns.append(f"  accent bar (y={sy:.2f}) overlaps title bottom (y={title_bottom:.2f})")

        if slide_warns:
            print(f"\n[slide {i}]")
            for w in slide_warns: print(w)
            warnings.extend(slide_warns)

    print()
    if warnings:
        print(f"Found {len(warnings)} issue(s).")
        return False
    else:
        print("All checks passed.")
        return True


def main():
    ap = argparse.ArgumentParser(description='Wuren PPT visual QA')
    ap.add_argument('pptx', help='PPTX file')
    args = ap.parse_args()
    sys.exit(0 if check_deck(args.pptx) else 1)


if __name__ == '__main__':
    main()
